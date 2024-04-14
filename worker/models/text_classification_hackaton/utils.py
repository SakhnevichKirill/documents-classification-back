from api.s3 import s3
from PyPDF2 import PdfReader
from docx import Document
import openpyxl

class FileProcessor:
    def __init__(self):
        self.s3_client = s3

    async def process_file(self, filename: str, file_id: str):
        # Определение типа файла по расширению
        extension = filename.split('.')[-1].lower()
        if extension not in ['pdf', 'docx', 'xlsx', "txt"]:
            raise Exception(f"Unexpected file type: {extension}")

        # Скачивание файла
        file_data = await self.s3_client.download_file(file_id)

        # Обработка файла в зависимости от расширения
        if extension == 'pdf':
            return self.extract_text_from_pdf(file_data)
        elif extension == 'docx':
            return self.extract_text_from_docx(file_data)
        elif extension == 'xlsx':
            return self.extract_text_from_xlsx(file_data)
        elif extension == 'txt':
            return self.extract_text_from_txt(file_data)

    def extract_text_from_pdf(self, file_data):
        with open('temp_file.pdf', 'wb') as f:
            f.write(file_data)
        reader = PdfReader('temp_file.pdf')
        text = ""
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:  # Проверяем, что текст был извлечен
                text += page_text.replace('\n', ' ') 
        return text

    def extract_text_from_docx(self, file_data):
        with open('temp_file.docx', 'wb') as f:
            f.write(file_data)
        doc = Document('temp_file.docx')
        text = "\n".join([para.text for para in doc.paragraphs])
        return text

    def extract_text_from_xlsx(self, file_data):
        with open('temp_file.xlsx', 'wb') as f:
            f.write(file_data)
        workbook = openpyxl.load_workbook('temp_file.xlsx')
        sheet = workbook.active
        text = "\n".join([str(cell.value) for row in sheet for cell in row if cell.value is not None])
        return text
    
    def extract_text_from_txt(self, file_data):
        return file_data.decode('utf-8')